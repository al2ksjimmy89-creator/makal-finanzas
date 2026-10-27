"""Importa los PyG mensuales de Excel (formato Escuela de Policía San Pablo) a la base de datos.

Uso desde la app: Configuración → Importar Excel.
Uso por consola:  python importar_excel.py "archivo.xlsx" 2026-10 "Escuela de Policía San Pablo"

Volver a importar el mismo mes reemplaza lo que había de ese mes para ese cliente.
"""
import sys
import unicodedata
import warnings
from calendar import monthrange
from datetime import date, datetime

import openpyxl

import db
from calculos import dia_corte, pyg

warnings.filterwarnings("ignore")

GRUPOS_MP = {"vegetales": "Vegetales", "secos": "Abastos", "abastos": "Abastos", "proteinas": "Proteínas",
             "desechables": "Desechables", "pulpas": "Pulpas", "pan": "Pan"}

# texto que aparece en "Observación" -> (proveedor, grupo habitual)
PROVEEDORES_CONOCIDOS = [
    ("virgin", "Virgiña Quilumbaqui", "Vegetales"), ("veg costa", "Veg Costa", "Vegetales"),
    ("avitalsa", "AVITALSA", "Proteínas"), ("janeth", "Janeth Flores", "Proteínas"),
    ("jeaneth", "Janeth Flores", "Proteínas"), ("normita", "Normita", "Proteínas"),
    ("iluman", "Tercena Iluman", "Proteínas"), ("tuti", "Tuti", "Proteínas"), ("yaguachi", "Yaguachi", "Proteínas"),
    ("titan", "TITAN", "Proteínas"),
    ("hormiguita", "Hormiguita", "Abastos"), ("horiguita", "Hormiguita", "Abastos"),
    ("hernandez", "Comercial Hernandez", "Abastos"), ("yepez", "Abastos Yepez", "Abastos"),
    ("san jose", "San Jose", "Abastos"), ("santa anita", "Santa Anita", "Abastos"),
    ("supermaxi", "Supermaxi", "Abastos"), ("borja rosa", "Rosa Borja", "Desechables"),
    ("farinango", "Oswaldo Farinango", "Pan"), ("fran fruit", "Fran Fruit", "Pulpas"),
]

HOJAS_SIMPLES = {"mantenimiento": "Repuestos y mantenimiento", "menaje": "Menaje y equipos",
                 "gastos varios": "Gastos varios / administrativos"}


def norm(v):
    if v is None:
        return ""
    t = unicodedata.normalize("NFKD", str(v)).encode("ascii", "ignore").decode().lower()
    return " ".join(t.replace(".", " ").split())


def num(v):
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def es_fecha(v):
    return isinstance(v, (datetime, date))


def hoja(wb, *nombres):
    for ws in wb.worksheets:
        if any(norm(ws.title).startswith(n) for n in nombres):
            return ws
    return None


def buscar(ws, condicion, col=None):
    for fila in ws.iter_rows():
        for c in fila:
            if (col is None or c.column == col) and condicion(norm(c.value)):
                return c.row, c.column
    return None, None


def proveedor_de(texto):
    t = norm(texto)
    for clave, nombre, grupo in PROVEEDORES_CONOCIDOS:
        if clave in t:
            return nombre, grupo
    return None, None


# ---------------------------------------------------------------- lectura de hojas

def leer_venta(ws, periodo):
    hr, _ = buscar(ws, lambda t: t == "fecha", col=1)
    servicios = {}
    for c in (2, 3, 4):
        nombre = norm(ws.cell(hr - 1, c).value)
        if nombre:
            servicios[c] = nombre.capitalize()
    precios = {}
    for fila in ws.iter_rows(min_col=10, max_col=12):
        t = norm(fila[0].value).rstrip("s")
        for nombre in servicios.values():
            if t == norm(nombre).rstrip("s") and isinstance(fila[2].value, (int, float)):
                precios[nombre] = float(fila[2].value)
    _, col_transf = buscar(ws, lambda t: "transferencia" in t and "pago" in t)

    raciones, transferencias = [], []
    r = hr + 1
    while r <= ws.max_row and norm(ws.cell(r, 1).value) != "total":
        f = ws.cell(r, 1).value
        if es_fecha(f):
            for c, serv in servicios.items():
                cant = num(ws.cell(r, c).value)
                if cant:
                    raciones.append((f.date(), serv, cant))
            if col_transf and num(ws.cell(r, col_transf).value):
                transferencias.append((f.date(), num(ws.cell(r, col_transf).value)))
        r += 1
    # a veces el total de transferencias se escribió a mano en la fila TOTAL
    if col_transf:
        total_fila = num(ws.cell(r, col_transf).value)
        diferencia = round(total_fila - sum(v for _, v in transferencias), 2)
        if diferencia > 0:
            transferencias.append((None, diferencia))

    instructores = []
    hn, _ = buscar(ws, lambda t: t == "nombre", col=1)
    if hn:
        r = hn + 1
        while r <= ws.max_row and "total" not in norm(ws.cell(r, 1).value) + norm(ws.cell(r, 2).value):
            nombre = ws.cell(r, 1).value
            if nombre:
                instructores.append((str(nombre).strip().title(), ws.cell(r, 2).value, num(ws.cell(r, 3).value)))
            r += 1

    he, ce = buscar(ws, lambda t: t == "escuela", col=9)
    escuela = num(ws.cell(he, 13).value) if he else 0.0
    return {"servicios": list(servicios.values()), "precios": precios, "raciones": raciones,
            "transferencias": transferencias, "instructores": instructores, "escuela": escuela}


def leer_materia_prima(ws):
    filas = []
    hr = None
    for fila in ws.iter_rows(max_row=8):
        grupos = {c.column: GRUPOS_MP[norm(c.value)] for c in fila if norm(c.value) in GRUPOS_MP}
        if len(grupos) >= 3:
            hr = fila[0].row
            break
    r = hr + 2
    while r <= ws.max_row and "total" not in norm(ws.cell(r, 2).value):
        f = ws.cell(r, 2).value
        if es_fecha(f):
            for c, grupo in grupos.items():
                if grupo == "Pan":
                    cant, pu = num(ws.cell(r, c).value), num(ws.cell(r, c + 1).value)
                    monto = num(ws.cell(r, c + 3).value) or cant * pu
                    if monto:
                        filas.append(dict(fecha=f.date(), grupo="Pan", monto=monto, cantidad=cant, precio_unit=pu,
                                          detalle=ws.cell(r, c + 2).value, n_factura=None, pago=None))
                else:
                    monto = num(ws.cell(r, c).value)
                    if monto:
                        filas.append(dict(fecha=f.date(), grupo=grupo, monto=monto, cantidad=None, precio_unit=None,
                                          detalle=ws.cell(r, c + 2).value, n_factura=ws.cell(r, c + 1).value,
                                          pago=ws.cell(r, c + 3).value))
        r += 1
    return filas


def leer_pan_antiguo(ws):
    """Formato marzo-abril: hoja 'Pan' con pan normal, especial y jugos."""
    filas = []
    hp, _ = buscar(ws, lambda t: t.startswith("cost"), col=3)
    precios = {4: num(ws.cell(hp, 4).value), 5: num(ws.cell(hp, 5).value)}
    hr, _ = buscar(ws, lambda t: t == "fecha", col=3)
    for r in range(hr + 1, hp):
        f = ws.cell(r, 3).value
        if not es_fecha(f):
            continue
        for c, tipo in ((4, "Pan normal"), (5, "Pan especial")):
            cant = num(ws.cell(r, c).value)
            if cant:
                filas.append(dict(fecha=f.date(), grupo="Pan", monto=cant * precios[c], cantidad=cant,
                                  precio_unit=precios[c], detalle=tipo, n_factura=None, pago=None))
    extra = num(ws.cell(hp + 1, 6).value)
    if extra:
        filas.append(dict(fecha=None, grupo="Pan", monto=extra, cantidad=None, precio_unit=None,
                          detalle=str(ws.cell(hp + 1, 7).value or "Otros").title(), n_factura=None, pago=None))
    return filas


def leer_glp(ws):
    hr, cf = buscar(ws, lambda t: t == "fecha")
    hp, _ = buscar(ws, lambda t: "unit" in t, col=cf)
    precio = num(ws.cell(hp, cf + 1).value) if hp else 3.5
    filas = []
    r = hr + 1
    while r <= ws.max_row and "total" not in norm(ws.cell(r, cf).value):
        f, cant = ws.cell(r, cf).value, num(ws.cell(r, cf + 1).value)
        if es_fecha(f) and cant:
            filas.append(dict(fecha=f.date(), cantidad=cant, precio_unit=precio, monto=cant * precio))
        r += 1
    return filas


def leer_hoja_simple(ws):
    hr, cv = buscar(ws, lambda t: t == "valor")
    filas = []
    r = hr + 1
    while r <= ws.max_row and "total" not in norm(ws.cell(r, cv - 1).value):
        f, monto = ws.cell(r, cv - 1).value, num(ws.cell(r, cv).value)
        if monto:
            filas.append(dict(fecha=f.date() if es_fecha(f) else None, monto=monto,
                              detalle=ws.cell(r, cv + 1).value, n_factura=ws.cell(r, cv + 2).value,
                              pago=ws.cell(r, cv + 3).value))
        r += 1
    return filas


def leer_mano_obra(ws):
    hr, _ = buscar(ws, lambda t: t == "nombre")
    cols = {norm(c.value): c.column for c in ws[hr] if c.value}

    def col(*claves):
        for k, c in cols.items():
            if any(k.startswith(clave) for clave in claves):
                return c
        return None

    c_cargo, c_ced, c_nom, c_sue = col("cargo"), col("cedula"), col("nombre"), col("sueldo")
    c_dias, c_bono, c_ing, c_sal = col("dias lab"), col("bono"), col("fecha ingreso"), col("fecha salida")
    c_desc, c_total = col("descuento"), col("total")
    base = 30.0
    hb, cb = buscar(ws, lambda t: t == "dias base")
    if hb:
        v = ws.cell(hb, cb + 1).value
        base = num(v) or num(ws.cell(hb + 1, cb).value) or 30.0
    empleados = []
    r = hr + 1
    while r <= ws.max_row and "total" not in norm(ws.cell(r, 2).value) + norm(ws.cell(r, 3).value):
        nombre, sueldo = ws.cell(r, c_nom).value, num(ws.cell(r, c_sue).value)
        if nombre and sueldo:
            d = ws.cell(r, c_desc) if c_desc else None
            empleados.append(dict(
                nombre=" ".join(str(nombre).split()), cargo=str(ws.cell(r, c_cargo).value or "").strip(),
                cedula=str(ws.cell(r, c_ced).value or "").strip() or None, sueldo=sueldo,
                dias=num(ws.cell(r, c_dias).value), bono=num(ws.cell(r, c_bono).value) if c_bono else 0.0,
                ingreso=ws.cell(r, c_ing).value, salida=ws.cell(r, c_sal).value,
                descuento=num(d.value) if d else 0.0,
                nota_descuento=d.comment.text.replace("Admin:", "").strip() if d is not None and d.comment else "",
                total=num(ws.cell(r, c_total).value) if c_total else 0.0))
        r += 1
    return empleados, base


def leer_totales_pyg(wb):
    ws = hoja(wb, "pyg")
    etiquetas = {"ventas": ("total ventas",), "mp": ("total materia prima", "total 511"),
                 "mo": ("total mano de obra", "total 513"),
                 "indirectos": ("total costos indirectos", "total 52 costos indirectos"),
                 "utilidad": ("utilidad final",)}
    res = {}
    for clave, textos in etiquetas.items():
        for fila in ws.iter_rows():
            for c in fila:
                t = norm(c.value)
                if any(x in t for x in textos):
                    res[clave] = num(ws.cell(c.row, 9).value)  # la última coincidencia es el total real
    return res


# Líneas del PyG que a veces se escribían a mano (sin hoja de detalle)
LINEAS_MANUALES_PYG = [
    ("directv", "Internet / TV"), ("internet", "Internet / TV"), ("plagas", "Control de plagas"),
    ("uniformes", "Uniformes e implementos"), ("arriendo", "Arriendo local"),
    ("suministros", "Gastos varios / administrativos"), ("transporte alimentos", "Transporte"),
    ("abono inversion", "Abono inversión inicial"),
]


def leer_lineas_manuales_pyg(ruta):
    """Valores numéricos escritos directamente en la hoja PyG (no fórmulas)."""
    wb = openpyxl.load_workbook(ruta, data_only=False)
    ws = hoja(wb, "pyg")
    filas = []
    for fila in ws.iter_rows():
        etiqueta = " ".join(norm(c.value) for c in fila if isinstance(c.value, str))
        v = ws.cell(fila[0].row, 9).value
        if isinstance(v, str) and v.startswith("="):
            # fórmulas tipo =54.5+17 también son valores a mano
            if not all(ch in "0123456789.+= " for ch in v):
                continue
            v = eval(v[1:])  # noqa: S307 - solo dígitos y '+'
        if not isinstance(v, (int, float)) or not v:
            continue
        for clave, categoria in LINEAS_MANUALES_PYG:
            if clave in etiqueta:
                filas.append((categoria, float(v), etiqueta.strip()))
                break
    return filas


# ---------------------------------------------------------------- escritura

def _categoria_id(con, nombre):
    return db.valor(con, "SELECT id FROM categorias WHERE nombre=?", (nombre,))


def asegurar_cliente(con, nombre, venta):
    cid = db.valor(con, "SELECT id FROM clientes WHERE nombre=?", (nombre,))
    if cid is None:
        cid = con.execute("INSERT INTO clientes (nombre, meta_utilidad, personas_nomina) VALUES (?,0.20,0) "
                          "RETURNING id", (nombre,)).fetchone()[0]
    for i, serv in enumerate(venta["servicios"]):
        precio = venta["precios"].get(serv, 0)
        con.execute("""INSERT INTO servicios_cliente (cliente_id, grupo, servicio, precio, orden)
                       VALUES (?, 'Aspirantes', ?, ?, ?)
                       ON CONFLICT(cliente_id, grupo, servicio) DO UPDATE SET precio=excluded.precio""",
                    (cid, serv, precio, i))
    return cid


def importar_archivo(con, ruta, nombre_cliente, periodo, hoy=None):
    """Importa un archivo y devuelve un informe con totales Excel vs app."""
    wb = openpyxl.load_workbook(ruta, data_only=True)
    cerrado = dia_corte(periodo, hoy) >= 30
    a, m = map(int, periodo.split("-"))
    ultimo_dia = date(a, m, monthrange(a, m)[1])

    venta = leer_venta(hoja(wb, "venta"), periodo)
    cid = asegurar_cliente(con, nombre_cliente, venta)

    for tabla in ("raciones", "ventas_mensuales", "ventas_extra", "gastos"):
        con.execute(f"DELETE FROM {tabla} WHERE cliente_id=? AND periodo=?", (cid, periodo))
    con.execute("""DELETE FROM novedades_rol WHERE periodo=? AND empleado_id IN
                   (SELECT id FROM empleados WHERE cliente_id=?)""", (periodo, cid))
    con.execute("""DELETE FROM descuentos WHERE periodo_inicio=? AND empleado_id IN
                   (SELECT id FROM empleados WHERE cliente_id=?)""", (periodo, cid))

    # Ventas
    for f, serv, cant in venta["raciones"]:
        con.execute("""INSERT INTO raciones (cliente_id, periodo, fecha, grupo, servicio, cantidad, precio)
                       VALUES (?,?,?,?,?,?,?)""", (cid, periodo, f.isoformat(), "Aspirantes", serv, cant,
                                                  venta["precios"].get(serv, 0)))
    for nombre, grado, valor in venta["instructores"]:
        con.execute("""INSERT INTO personas_mensuales (cliente_id, grupo, nombre, grado, valor_mensual)
                       VALUES (?, 'Instructores', ?, ?, 40) ON CONFLICT(cliente_id, nombre)
                       DO UPDATE SET grado=excluded.grado""", (cid, nombre, grado))
        if valor:
            pid = db.valor(con, "SELECT id FROM personas_mensuales WHERE cliente_id=? AND nombre=?", (cid, nombre))
            con.execute("INSERT INTO ventas_mensuales (cliente_id, periodo, persona_id, monto) VALUES (?,?,?,?)",
                        (cid, periodo, pid, valor))
    for f, valor in venta["transferencias"]:
        con.execute("""INSERT INTO ventas_extra (cliente_id, periodo, fecha, tipo, detalle, monto, forma_pago)
                       VALUES (?,?,?,'Personal externo','Transferencias (importado)',?, 'Transferencia')""",
                    (cid, periodo, (f or ultimo_dia).isoformat(), valor))
    if venta["escuela"]:
        con.execute("""INSERT INTO ventas_extra (cliente_id, periodo, fecha, tipo, detalle, monto)
                       VALUES (?,?,?,'Evento extra','Eventos pagados por la Escuela (importado)',?)""",
                    (cid, periodo, ultimo_dia.isoformat(), venta["escuela"]))

    # Gastos
    def pagado(marca):
        return 1 if cerrado or (marca and "☑" in str(marca)) else 0

    def insertar_gasto(fecha, categoria, monto, detalle=None, proveedor=None, cantidad=None, precio_unit=None,
                       n_factura=None, marca=None):
        prov_id = None
        if proveedor:
            prov_id = db.obtener_o_crear_proveedor(con, proveedor[0], _categoria_id(con, proveedor[1]))
        con.execute("""INSERT INTO gastos (cliente_id, periodo, fecha, categoria_id, proveedor_id, detalle,
                       cantidad, precio_unit, monto, n_factura, pagado) VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                    (cid, periodo, (fecha or ultimo_dia).isoformat(), _categoria_id(con, categoria), prov_id,
                     str(detalle).strip() if detalle else None, cantidad or None, precio_unit or None, monto,
                     str(n_factura) if n_factura else None, pagado(marca)))

    mp = leer_materia_prima(hoja(wb, "materia prima"))
    ws_pan = hoja(wb, "pan")
    if ws_pan:
        mp += leer_pan_antiguo(ws_pan)
    for g in mp:
        prov = proveedor_de(g["detalle"])
        detalle = g["detalle"]
        if prov[0] and norm(detalle) in (norm(prov[0]), *[k for k, n, _ in PROVEEDORES_CONOCIDOS if n == prov[0]]):
            detalle = None  # la observación era solo el nombre del proveedor
        insertar_gasto(g["fecha"], g["grupo"], g["monto"], detalle, prov if prov[0] else None,
                       g["cantidad"], g["precio_unit"], g["n_factura"], g["pago"])
    for g in leer_glp(hoja(wb, "glp")):
        insertar_gasto(g["fecha"], "Combustible / Gas (GLP)", g["monto"], "Tanques 15 kg", None,
                       g["cantidad"], g["precio_unit"])
    for clave, categoria in HOJAS_SIMPLES.items():
        ws = hoja(wb, clave)
        if ws:
            for g in leer_hoja_simple(ws):
                insertar_gasto(g["fecha"], categoria, g["monto"], g["detalle"], None, None, None,
                               g["n_factura"], g["pago"])
    for categoria, monto, etiqueta in leer_lineas_manuales_pyg(ruta):
        insertar_gasto(None, categoria, monto, f"Valor escrito en el PyG: {etiqueta.capitalize()}", marca="☑")

    # Mano de obra
    empleados, base = leer_mano_obra(hoja(wb, "mano obra"))
    ids_empleados = []
    for e in empleados:
        eid = db.valor(con, "SELECT id FROM empleados WHERE cliente_id=? AND lower(nombre)=lower(?) AND cargo=?",
                       (cid, e["nombre"], e["cargo"]))
        ingreso = db.a_texto_fecha(e["ingreso"])
        salida = db.a_texto_fecha(e["salida"])
        if eid is None:
            eid = con.execute("""INSERT INTO empleados (cliente_id, cedula, nombre, cargo, sueldo, bono,
                                 tipo_contrato, fecha_ingreso, fecha_salida) VALUES (?,?,?,?,?,?,?,?,?)
                                 RETURNING id""",
                              (cid, e["cedula"], e["nombre"], e["cargo"], e["sueldo"], e["bono"],
                               "Prestación de servicios", ingreso, salida)).fetchone()[0]
        else:
            con.execute("""UPDATE empleados SET cedula=COALESCE(?,cedula), sueldo=?, bono=?, fecha_ingreso=?,
                           fecha_salida=? WHERE id=?""", (e["cedula"], e["sueldo"], e["bono"], ingreso, salida, eid))
        ids_empleados.append(eid)
        if cerrado:
            # mes cerrado: se congela tal como figura en el Excel (días convertidos a base 30).
            # Si el total pagado se escribió a mano, se ajustan los días para respetar lo pagado.
            dias30 = e["dias"] * 30 / base
            if e["total"] and abs(e["sueldo"] * dias30 / 30 + e["bono"] - e["total"]) > 0.01 and not e["descuento"]:
                dias30 = (e["total"] - e["bono"]) * 30 / e["sueldo"]
            con.execute("""INSERT INTO novedades_rol (empleado_id, periodo, dias, sueldo, bono, observacion)
                           VALUES (?,?,?,?,?,'Importado de Excel') ON CONFLICT(empleado_id, periodo)
                           DO UPDATE SET dias=dias+excluded.dias, bono=bono+excluded.bono""",
                        (eid, periodo, round(dias30, 4), e["sueldo"], e["bono"]))
        if e["descuento"]:
            nota = e["nota_descuento"]
            tipo = "Recupera gasto de la empresa" if any(x in norm(nota) for x in ("dotacion", "calzado")) \
                else "Adelanto de sueldo"
            con.execute("""INSERT INTO descuentos (empleado_id, periodo_inicio, tipo, motivo, monto_total, cuotas)
                           VALUES (?,?,?,?,?,1)""", (eid, periodo, tipo, nota or "Importado de Excel", e["descuento"]))
    if cerrado:
        con.execute("INSERT INTO roles_cerrados VALUES (?,?,?) ON CONFLICT (cliente_id, periodo) DO NOTHING",
                    (cid, periodo, db.hoy().isoformat()))
    else:
        con.execute("DELETE FROM roles_cerrados WHERE cliente_id=? AND periodo=?", (cid, periodo))
    con.commit()
    db.aplicar_cuadro_proveedores(con)

    excel = leer_totales_pyg(wb)
    r = pyg(con, cid, periodo, dia_corte(periodo, hoy))
    app = {"ventas": r["total_ventas"], "mp": r["total_mp"], "mo": r["total_mo"],
           "indirectos": r["indirectos"], "utilidad": r["utilidad_neta"]}
    return {"cliente_id": cid, "periodo": periodo, "excel": excel, "app": app, "cerrado": cerrado,
            "empleados": ids_empleados}


def ajustar_estado_empleados(con, empleados_por_periodo):
    """Después de importar varios meses: quien no figura en el último mes queda inactivo,
    con fecha de salida al final del último mes en que aparece (si no tenía una)."""
    ultimo = max(empleados_por_periodo)
    ultima_aparicion = {}
    for periodo in sorted(empleados_por_periodo):
        for eid in empleados_por_periodo[periodo]:
            ultima_aparicion[eid] = periodo
    for eid, periodo in ultima_aparicion.items():
        if periodo == ultimo:
            con.execute("UPDATE empleados SET activo=1 WHERE id=?", (eid,))
        else:
            a, m = map(int, periodo.split("-"))
            con.execute("UPDATE empleados SET activo=0, fecha_salida=COALESCE(fecha_salida, ?) WHERE id=?",
                        (date(a, m, monthrange(a, m)[1]).isoformat(), eid))
    con.commit()


if __name__ == "__main__":
    con = db.inicializar()
    informe = importar_archivo(con, sys.argv[1], sys.argv[3] if len(sys.argv) > 3 else "Escuela de Policía San Pablo",
                               sys.argv[2])
    print(informe)
